TARGET_EQUATION = "C127_3434"
EXPECTED_CORRIDORS = 3
EQUATIONS_PER_CORRIDOR = 55
EXPECTED_TOTAL_EQUATIONS = 165


def validate_engine_result(engine_result):
    corridors = engine_result.get("corridors", [])

    if len(corridors) != EXPECTED_CORRIDORS:
        raise ValueError(
            f"QC FAILED: expected {EXPECTED_CORRIDORS} corridors, "
            f"received {len(corridors)}"
        )

    total_equations = 0

    for index, corridor in enumerate(corridors, start=1):
        equations = corridor.get("full_55_equations", [])

        if len(equations) != EQUATIONS_PER_CORRIDOR:
            raise ValueError(
                f"QC FAILED: corridor {index} contains "
                f"{len(equations)} equations; expected "
                f"{EQUATIONS_PER_CORRIDOR}"
            )

        if equations[-1] != TARGET_EQUATION:
            raise ValueError(
                f"QC FAILED: corridor {index} ends at "
                f"{equations[-1]}, not {TARGET_EQUATION}"
            )

        total_equations += len(equations)

    if total_equations != EXPECTED_TOTAL_EQUATIONS:
        raise ValueError(
            f"QC FAILED: received {total_equations} total equations; "
            f"expected {EXPECTED_TOTAL_EQUATIONS}"
        )

    return {
        "status": "qc_passed",
        "corridor_count": len(corridors),
        "equations_per_corridor": EQUATIONS_PER_CORRIDOR,
        "total_equations": total_equations,
        "final_equation": TARGET_EQUATION,
    }
from pathlib import Path
from docx import Document


EXPECTED_TEMPLATE_PARAGRAPHS = 508

PERSONAL_SLOTS = {
    "opening_customer_name": 3,
    "opening_numeric_name": 8,
    "manuscript_number": 28,
    "calculation_customer_name": 34,
    "name_value": 38,
    "birth_month_day": 42,
    "calculation_numeric_name": 48,
}

CORRIDOR_A_SLOTS = (
    list(range(72, 83))
    + list(range(95, 106))
    + list(range(115, 136, 2))
    + list(range(144, 155))
    + list(range(164, 185, 2))
)

CORRIDOR_B_SLOTS = (
    list(range(214, 225))
    + list(range(235, 246))
    + list(range(255, 266))
    + list(range(275, 286))
    + list(range(294, 305))
)

CORRIDOR_C_SLOTS = (
    list(range(333, 344))
    + list(range(354, 365))
    + list(range(374, 385))
    + list(range(394, 405))
    + list(range(413, 424))
)

CORRIDOR_SLOTS = (
    CORRIDOR_A_SLOTS,
    CORRIDOR_B_SLOTS,
    CORRIDOR_C_SLOTS,
)

PROTECTED_ANCHORS = {
    465: "The Convergence Argument",
    482: "Preserving Your Numeromancy Manuscript",
    495: "Archive Annotation",
}

EXPECTED_TEMPLATE_VALUES = {
    3: "Emma Reynolds",
    8: "180",
    28: "Manuscript No. 0002",
    34: "EMMA REYNOLDS",
    38: "144",
    42: "30 + 6",
    48: "180",
}


def _replace_paragraph_text(paragraph, value):
    value = str(value)

    if paragraph.runs:
        paragraph.runs[0].text = value

        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _validate_template(document):
    if len(document.paragraphs) != EXPECTED_TEMPLATE_PARAGRAPHS:
        raise ValueError(
            "MANUSCRIPT QC FAILED: template must contain "
            f"exactly {EXPECTED_TEMPLATE_PARAGRAPHS} paragraphs; "
            f"found {len(document.paragraphs)}"
        )

    for paragraph_index, expected_text in EXPECTED_TEMPLATE_VALUES.items():
        actual_text = document.paragraphs[paragraph_index].text.strip()

        if actual_text != expected_text:
            raise ValueError(
                "MANUSCRIPT QC FAILED: template fingerprint mismatch "
                f"at paragraph {paragraph_index}; "
                f"expected {expected_text!r}, found {actual_text!r}"
            )

    for paragraph_index, expected_text in PROTECTED_ANCHORS.items():
        actual_text = document.paragraphs[paragraph_index].text.strip()

        if expected_text not in actual_text:
            raise ValueError(
                "MANUSCRIPT QC FAILED: protected anchor missing "
                f"at paragraph {paragraph_index}; "
                f"expected text containing {expected_text!r}"
            )


def render_manuscript(
    template_path,
    output_path,
    customer_name,
    dob,
    manuscript_number,
    engine_result,
):
    validate_engine_result(engine_result)

    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.is_file():
        raise FileNotFoundError(
            f"MANUSCRIPT QC FAILED: template not found: {template_path}"
        )

    if template_path.suffix.lower() != ".docx":
        raise ValueError(
            "MANUSCRIPT QC FAILED: template must be a DOCX file"
        )

    if output_path.suffix.lower() != ".docx":
        raise ValueError(
            "MANUSCRIPT QC FAILED: output must be a DOCX file"
        )

    if template_path.resolve() == output_path.resolve():
        raise ValueError(
            "MANUSCRIPT QC FAILED: output path must not overwrite "
            "the master template"
        )

    document = Document(str(template_path))
    _validate_template(document)

    corridors = engine_result["corridors"]

    if len(CORRIDOR_A_SLOTS) != EQUATIONS_PER_CORRIDOR:
        raise ValueError(
            "MANUSCRIPT QC FAILED: Corridor A slot map is not 55 positions"
        )

    if len(CORRIDOR_B_SLOTS) != EQUATIONS_PER_CORRIDOR:
        raise ValueError(
            "MANUSCRIPT QC FAILED: Corridor B slot map is not 55 positions"
        )

    if len(CORRIDOR_C_SLOTS) != EQUATIONS_PER_CORRIDOR:
        raise ValueError(
            "MANUSCRIPT QC FAILED: Corridor C slot map is not 55 positions"
        )

    name_value = sum(
        (ord(character.lower()) - 96)
        for character in customer_name
        if character.isalpha()
    )

    try:
        _, month, day = map(int, dob.split("-"))
    except (AttributeError, TypeError, ValueError):
        raise ValueError(
            "MANUSCRIPT QC FAILED: dob must use YYYY-MM-DD format"
        )

    calculated_numeric_name = name_value + month + day
    engine_numeric_name = engine_result.get("numeric_name")

    if calculated_numeric_name != engine_numeric_name:
        raise ValueError(
            "MANUSCRIPT QC FAILED: calculated Numeric Name "
            f"{calculated_numeric_name} does not match engine Numeric Name "
            f"{engine_numeric_name}"
        )

    manuscript_number = str(manuscript_number).strip()

    if not manuscript_number:
        raise ValueError(
            "MANUSCRIPT QC FAILED: manuscript number is required"
        )

    personal_values = {
        PERSONAL_SLOTS["opening_customer_name"]: customer_name,
        PERSONAL_SLOTS["opening_numeric_name"]: calculated_numeric_name,
        PERSONAL_SLOTS["manuscript_number"]:
            f"Manuscript No. {manuscript_number.zfill(4)}",
        PERSONAL_SLOTS["calculation_customer_name"]:
            customer_name.upper(),
        PERSONAL_SLOTS["name_value"]: name_value,
        PERSONAL_SLOTS["birth_month_day"]: f"{month} + {day}",
        PERSONAL_SLOTS["calculation_numeric_name"]:
            calculated_numeric_name,
    }

    for paragraph_index, value in personal_values.items():
        _replace_paragraph_text(
            document.paragraphs[paragraph_index],
            value,
        )

    equations_written = 0

    for corridor_index, paragraph_slots in enumerate(CORRIDOR_SLOTS):
        corridor = corridors[corridor_index]
        equations = corridor.get("full_55_equations", [])

        if not isinstance(equations, list):
            raise ValueError(
                "MANUSCRIPT QC FAILED: corridor "
                f"{corridor_index + 1} does not contain an equation list"
            )

        if len(equations) != EQUATIONS_PER_CORRIDOR:
            raise ValueError(
                "MANUSCRIPT QC FAILED: corridor "
                f"{corridor_index + 1} contains {len(equations)} equations; "
                f"expected {EQUATIONS_PER_CORRIDOR}"
            )

        for paragraph_index, equation in zip(
            paragraph_slots,
            equations,
        ):
            _replace_paragraph_text(
                document.paragraphs[paragraph_index],
                equation,
            )
            equations_written += 1

    if equations_written != EXPECTED_TOTAL_EQUATIONS:
        raise ValueError(
            "MANUSCRIPT QC FAILED: wrote "
            f"{equations_written} equations; "
            f"expected {EXPECTED_TOTAL_EQUATIONS}"
        )

    for final_slot in (184, 304, 423):
        if document.paragraphs[final_slot].text.strip() != TARGET_EQUATION:
            raise ValueError(
                "MANUSCRIPT QC FAILED: final equation missing "
                f"from paragraph {final_slot}"
            )

    for paragraph_index, expected_text in PROTECTED_ANCHORS.items():
        actual_text = document.paragraphs[paragraph_index].text.strip()

        if expected_text not in actual_text:
            raise ValueError(
                "MANUSCRIPT QC FAILED: protected content changed "
                f"at paragraph {paragraph_index}"
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))

    saved_document = Document(str(output_path))

    if len(saved_document.paragraphs) != EXPECTED_TEMPLATE_PARAGRAPHS:
        output_path.unlink(missing_ok=True)

        raise ValueError(
            "MANUSCRIPT QC FAILED: saved manuscript structure changed"
        )

    for paragraph_index, expected_value in personal_values.items():
        actual_value = saved_document.paragraphs[
            paragraph_index
        ].text.strip()

        if actual_value != str(expected_value):
            output_path.unlink(missing_ok=True)

            raise ValueError(
                "MANUSCRIPT QC FAILED: personal value verification "
                f"failed at paragraph {paragraph_index}"
            )

    for corridor_index, paragraph_slots in enumerate(CORRIDOR_SLOTS):
        corridor = corridors[corridor_index]
        equations = corridor.get("full_55_equations", [])

        for paragraph_index, expected_equation in zip(
            paragraph_slots,
            equations,
        ):
            actual_equation = saved_document.paragraphs[
                paragraph_index
            ].text.strip()

            if actual_equation != str(expected_equation):
                output_path.unlink(missing_ok=True)

                raise ValueError(
                    "MANUSCRIPT QC FAILED: equation verification "
                    f"failed at paragraph {paragraph_index}"
                )

    for paragraph_index, expected_text in PROTECTED_ANCHORS.items():
        actual_text = saved_document.paragraphs[
            paragraph_index
        ].text.strip()

        if expected_text not in actual_text:
            output_path.unlink(missing_ok=True)

            raise ValueError(
                "MANUSCRIPT QC FAILED: protected content changed "
                "after save"
            )

    return {
        "status": "manuscript_render_passed",
        "customer_name": customer_name,
        "numeric_name": calculated_numeric_name,
        "manuscript_number": manuscript_number.zfill(4),
        "paragraph_count": len(saved_document.paragraphs),
        "corridor_count": len(corridors),
        "equations_written": equations_written,
        "final_equation": TARGET_EQUATION,
        "output_path": str(output_path),
    }
